from django.shortcuts import render,redirect,get_object_or_404
from django.contrib.auth.decorators import login_required


from donor.models import Donor
from patient.models import Patient
from blood.models import BloodStock, BloodRequest, DonationRequest
from accounts.models import ContactMessage
from django.db.models import Q
import json
from django.utils.dateparse import parse_date

from django.db.models import Count
from django.db.models.functions import TruncMonth

from django.http import HttpResponse
from django.core.mail import send_mail
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph,Spacer
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from django.utils import timezone
from reportlab.lib.pagesizes import A4
from reportlab.lib.enums import TA_CENTER
from django.contrib import messages

from accounts.models import Notification

#Excel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment








@login_required
def admin_dashboard(request):

    total_donors = Donor.objects.count()

    total_patients = Patient.objects.count()

    total_blood_groups = BloodStock.objects.count()

    total_blood_units = sum(
        stock.units for stock in BloodStock.objects.all()
    )

    pending_blood_requests = BloodRequest.objects.filter(
        status="Pending"
    ).count()

    pending_donation_requests = DonationRequest.objects.filter(
        status="Pending"
    ).count()

    recent_blood_requests = BloodRequest.objects.order_by("-id")[:5]

    recent_donation_requests = DonationRequest.objects.order_by("-id")[:5]

    context = {

         "recent_blood_requests": recent_blood_requests,

        "total_donors": total_donors,

        "total_patients": total_patients,

        "total_blood_groups": total_blood_groups,

        "total_blood_units": total_blood_units,

        "pending_blood_requests": pending_blood_requests,

        "pending_donation_requests": pending_donation_requests,

        "recent_blood_requests": recent_blood_requests,

        "recent_donation_requests": recent_donation_requests,

    }

    return render(
        request,
        "adminpanel/dashboard.html",
        context
    )

def blood_requests(request):
    requests = BloodRequest.objects.select_related(
        "patient",
        "patient__user"
    ).order_by("-id")

    return render(
        request,
        "adminpanel/blood_requests.html",
        {
            "requests": requests
        }
    )

@login_required
def approve_blood_request(request, id):

    blood_request = BloodRequest.objects.get(id=id)

    blood_request.status = "Approved"
    blood_request.save()

    # Patient Notification
    Notification.objects.create(
        user=blood_request.patient.user,
        title="Blood Request Approved",
        message=f"Congratulations! Your request for {blood_request.blood_group} blood has been approved.",
        notification_type="blood"
    )

    messages.success(request, "Blood Request Approved.")

    return redirect("blood_requests")

@login_required
def reject_blood_request(request, id):

    blood_request = BloodRequest.objects.get(id=id)

    blood_request.status = "Rejected"
    blood_request.save()

    # Patient Notification
    Notification.objects.create(
        user=blood_request.patient.user,
        title="Blood Request Rejected",
        message=f"Sorry! Your request for {blood_request.blood_group} blood has been rejected.",
        notification_type="blood"
    )

    messages.success(request, "Blood Request Rejected.")

    return redirect("blood_requests")      

def contact_messages(request):

    contact_messages = ContactMessage.objects.all().order_by("-created_at")

    return render(
    request,
    "adminpanel/contact messages.html",
    {"contact_messages": contact_messages}
)


def contact_message_detail(request, id):

    message = get_object_or_404(ContactMessage, id=id)
    message.is_read = True
    message.save(update_fields=["is_read"])

    return render(
        request,
        "adminpanel/contact_message_detail.html",
        {"message": message}
    )


@login_required
def reply_contact_message(request, id):
    message = get_object_or_404(ContactMessage, id=id)

    if request.method == "POST":
        reply_subject = request.POST.get("reply_subject", "").strip()
        reply_message = request.POST.get("reply_message", "").strip()

        if not reply_subject or not reply_message:
            messages.error(request, "Please enter both subject and message before sending a reply.")
            return redirect("contact_message_detail", id=message.id)

        original_message = (
            f"\n\n--- Original message from {message.name} ---\n"
            f"Subject: {message.subject}\n\n{message.message}"
        )

        final_email_body = (
            f"Hello {message.name},\n\n"
            f"Thank you for contacting LifeServe.\n\n"
            f"{reply_message}{original_message}"
        )

        send_mail(
            subject=reply_subject,
            message=final_email_body,
            from_email=None,
            recipient_list=[message.email],
            fail_silently=False,
        )

        messages.success(request, f"Reply sent successfully to {message.email}.")
        return redirect("contact_message_detail", id=message.id)

    return redirect("contact_message_detail", id=message.id)


def delete_contact_message(request, id):

    message = get_object_or_404(ContactMessage, id=id)

    message.delete()
    messages.success(request, "Message deleted successfully.")
    return redirect("contact_messages")






def manage_donors(request):

    search = request.GET.get("search")

    donors = Donor.objects.all().order_by("-id")

    if search:

        donors = donors.filter(user__first_name__icontains=search)

    context = {

        "donors": donors,

    }

    return render(
        request,
        "adminpanel/manage_donors.html",
        context
    )

def view_donor(request, id):
    donor = get_object_or_404(Donor, id=id)
    return render(request, "adminpanel/view_donor.html", {"donor": donor})


def edit_donor(request, id):
    donor = get_object_or_404(Donor, id=id)

    if request.method == "POST":
        donor.phone = request.POST.get("phone")
        donor.blood_group = request.POST.get("blood_group")
        donor.city = request.POST.get("city")
        donor.state = request.POST.get("state")
        donor.address = request.POST.get("address")

        donor.save()

        return redirect("manage_donors")

    return render(request, "adminpanel/edit_donor.html", {
        "donor": donor
    })
def delete_donor(request, id):
    donor = get_object_or_404(Donor, id=id)

    if request.method == "POST":
        donor.delete()
        messages.success(request, "Donor deleted successfully.")

    return redirect("manage_donors")



def manage_patients(request):
    search = request.GET.get("search", "")

    patients = Patient.objects.all().order_by("-id")

    if search:
        patients = patients.filter(
            Q(user__first_name__icontains=search) |
            Q(user__last_name__icontains=search) |
            Q(user__email__icontains=search) |
            Q(phone__icontains=search) |
            Q(blood_group__icontains=search) |
            Q(city__icontains=search)
        )

    context = {
        "patients": patients,
    }

    return render(request, "adminpanel/manage_patients.html", context)

def view_patient(request, id):
    patient = get_object_or_404(Patient, id=id)

    context = {
        "patient": patient,
    }

    return render(request, "adminpanel/view_patient.html", context)

def edit_patient(request, id):
    patient = get_object_or_404(Patient, id=id)

    if request.method == "POST":
        patient.phone = request.POST.get("phone")
        patient.blood_group = request.POST.get("blood_group")
        patient.city = request.POST.get("city")
        patient.state = request.POST.get("state")
        patient.address = request.POST.get("address")

        patient.save()

        return redirect("manage_patients")

    return render(request, "adminpanel/edit_patient.html", {
        "patient": patient
    })

def delete_patient(request, id):
    patient = get_object_or_404(Patient, id=id)

    if request.method == "POST":
        patient.delete()
        messages.success(request, "Patient deleted successfully.")

    return redirect("manage_patients")

def blood_stock(request):

    stocks = BloodStock.objects.all().order_by("blood_group")

    context = {
        "stocks": stocks
    }

    return render(request, "adminpanel/blood_stock.html", context)

def add_blood_stock(request):

        if request.method == "POST":

            blood_group = request.POST.get("blood_group")
            units = request.POST.get("units")

            # Same blood group already exists?
            if not BloodStock.objects.filter(blood_group=blood_group).exists():

                BloodStock.objects.create(
                    blood_group=blood_group,
                    units=units
                )

            return redirect("blood_stock")

        blood_groups = BloodStock.BLOOD_GROUPS

        return render(request, "adminpanel/add_blood_stock.html", {
            "blood_groups": blood_groups
        })


def edit_blood_stock(request, id):

    stock = get_object_or_404(BloodStock, id=id)

    if request.method == "POST":

        stock.units = request.POST.get("units")
        stock.save()

        return redirect("blood_stock")

    return render(request, "adminpanel/edit_blood_stock.html", {
        "stock": stock
    })


def delete_blood_stock(request, id):

    stock = get_object_or_404(BloodStock, id=id)

    if request.method == "POST":
        stock.delete()
        messages.success(request, "Blood stock deleted successfully.")

    return redirect("blood_stock")

def donation_requests(request):

    requests = DonationRequest.objects.select_related("donor", "donor__user").order_by("-id")

    context = {
        "requests": requests,
    }

    return render(request, "adminpanel/donation_requests.html", context)

def approve_donation_request(request, id):

    donation = get_object_or_404(DonationRequest, id=id)

    if donation.status == "Pending":

        donation.status = "Approved"
        donation.save()

        stock, created = BloodStock.objects.get_or_create(
            blood_group=donation.blood_group,
            defaults={"units": 0}
        )

        stock.units += donation.units
        stock.save()

        Notification.objects.create(
            user=donation.donor.user,
            title="🩸 Donation Request Approved",
            message=(
                f"Congratulations {donation.donor.user.first_name}! "
                "Your blood donation request has been approved. "
                "Thank you for saving lives ❤️."
            )
        )

        messages.success(request, "Donation request approved successfully.")

    return redirect("donation_requests")

def reject_donation_request(request, id):

    donation = get_object_or_404(DonationRequest, id=id)

    donation.status = "Rejected"
    donation.save()

    Notification.objects.create(
        user=donation.donor.user,
        title="❌ Donation Request Rejected",
        message=(
            f"Hello {donation.donor.user.first_name}, "
            "your blood donation request has been rejected. "
            "Please contact the blood bank for more information."
        )
    )

    messages.error(request, "Donation request rejected.")

    return redirect("donation_requests")





def reports(request):

    
    # Date Filter
    

    from_date = request.GET.get("from_date")
    to_date = request.GET.get("to_date")

    blood_requests = BloodRequest.objects.all()
    donation_requests = DonationRequest.objects.all()

    if from_date and to_date:

        blood_requests = blood_requests.filter(
        request_date__range=[
        parse_date(from_date),
        parse_date(to_date)
    ]
)

        donation_requests = donation_requests.filter(
        donation_date__range=[
        parse_date(from_date),
        parse_date(to_date)
    ]
)

    
    # Blood Stock
    

    blood_stock = BloodStock.objects.all().order_by("blood_group")

    blood_groups = []
    stock_units = []

    for stock in blood_stock:
        blood_groups.append(stock.blood_group)
        stock_units.append(stock.units)

    
    # Monthly Donations


    monthly_donations = (
    donation_requests
    .filter(status="Approved")
    .annotate(month=TruncMonth("donation_date"))
    .values("month")
    .annotate(total=Count("id"))
    .order_by("month")
)

    months = []
    donation_counts = []

    for item in monthly_donations:

        months.append(item["month"].strftime("%b"))
        donation_counts.append(item["total"])

    
    # Monthly Blood Requests
    

    monthly_requests = (
    blood_requests
    .annotate(month=TruncMonth("request_date"))
    .values("month")
    .annotate(total=Count("id"))
    .order_by("month")
)

    request_months = []
    request_counts = []

    for item in monthly_requests:

        request_months.append(item["month"].strftime("%b"))
        request_counts.append(item["total"])


    # Most Requested Blood Group
    

    most_requested = (
        BloodRequest.objects
        .values("blood_group")
        .annotate(total=Count("id"))
        .order_by("-total")
        .first()
    )


    # Most Active Donor
    

    most_active_donor = (
        DonationRequest.objects
        .filter(status="Approved")
        .values(
            "donor__user__first_name",
            "donor__user__last_name"
        )
        .annotate(total=Count("id"))
        .order_by("-total")
        .first()
    )

    
    # Low Stock
    

    low_stock = BloodStock.objects.filter(units__lt=10)

    
    # Context
    

    context = {

        "from_date": from_date,
        "to_date": to_date,

        "total_donors": Donor.objects.count(),
        "total_patients": Patient.objects.count(),

        "total_blood_requests": blood_requests.count(),
        "total_donation_requests": donation_requests.count(),

        "blood_stock": blood_stock,

        "blood_groups": json.dumps(blood_groups),
        "stock_units": json.dumps(stock_units),

        "approved_requests": blood_requests.filter(
            status="Approved"
        ).count(),

        "pending_requests": blood_requests.filter(
            status="Pending"
        ).count(),

        "rejected_requests": blood_requests.filter(
            status="Rejected"
        ).count(),

        "months": json.dumps(months),
        "donation_counts": json.dumps(donation_counts),

        "request_months": json.dumps(request_months),
        "request_counts": json.dumps(request_counts),

        "most_requested": most_requested,

        "most_active_donor": most_active_donor,

        "low_stock": low_stock,
    }

    return render(
        request,
        "adminpanel/reports.html",
        context
    )

def export_report_pdf(request):

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="LifeServe_Report.pdf"'

    doc = SimpleDocTemplate(
        response,
        pagesize=A4,
        rightMargin=25,
        leftMargin=25,
        topMargin=25,
        bottomMargin=25
    )

    styles = getSampleStyleSheet()

    title = styles["Title"]
    title.alignment = TA_CENTER

    heading = styles["Heading2"]

    normal = styles["BodyText"]

    elements = []

    
    # Title
    

    elements.append(
        Paragraph(
            "🩸 LifeServe Blood Bank Management System",
            title
        )
    )

    elements.append(
        Paragraph(
            "Administrative Report",
            heading
        )
    )

    elements.append(
        Paragraph(
            f"Generated on : {timezone.now().strftime('%d-%m-%Y %I:%M %p')}",
            normal
        )
    )

    elements.append(Spacer(1, 0.25 * inch))

    
    # Summary
    

    summary = [

        ["Metric", "Count"],

        ["Total Donors", Donor.objects.count()],

        ["Total Patients", Patient.objects.count()],

        ["Blood Requests", BloodRequest.objects.count()],

        ["Donation Requests", DonationRequest.objects.count()],

    ]

    summary_table = Table(summary, colWidths=[250, 150])

    summary_table.setStyle(TableStyle([

        ("BACKGROUND", (0,0), (-1,0), colors.darkred),

        ("TEXTCOLOR", (0,0), (-1,0), colors.white),

        ("GRID", (0,0), (-1,-1), 1, colors.grey),

        ("ALIGN", (0,0), (-1,-1), "CENTER"),

        ("BOTTOMPADDING", (0,0), (-1,0), 10),

        ("BACKGROUND", (0,1), (-1,-1), colors.whitesmoke),

    ]))

    elements.append(summary_table)

    elements.append(Spacer(1, 0.3 * inch))

    
    # Most Requested Blood Group
    

    most_requested = (

        BloodRequest.objects

        .values("blood_group")

        .annotate(total=Count("id"))

        .order_by("-total")

        .first()

    )

    if most_requested:

        elements.append(

            Paragraph(

                f"<b>Most Requested Blood Group:</b> "
                f"{most_requested['blood_group']} "
                f"({most_requested['total']} Requests)",

                normal

            )

        )

    
    # Most Active Donor


    donor = (

        DonationRequest.objects

        .filter(status="Approved")

        .values(
            "donor__user__first_name",
            "donor__user__last_name"
        )

        .annotate(total=Count("id"))

        .order_by("-total")

        .first()

    )

    if donor:

        elements.append(

            Paragraph(

                f"<b>Most Active Donor:</b> "
                f"{donor['donor__user__first_name']} "
                f"{donor['donor__user__last_name']} "
                f"({donor['total']} Donations)",

                normal

            )

        )

    elements.append(Spacer(1, 0.3 * inch))

    
    # Low Stock
    

    elements.append(
        Paragraph("Low Stock Alerts", heading)
    )

    low_stock = BloodStock.objects.filter(units__lt=10)

    if low_stock.exists():

        for stock in low_stock:

            elements.append(

                Paragraph(

                    f"• {stock.blood_group} : Only {stock.units} units remaining",

                    normal

                )

            )

    else:

        elements.append(

            Paragraph(

                "No low stock alerts.",

                normal

            )

        )

    elements.append(Spacer(1, 0.3 * inch))

    
    # Blood Stock Table
    

    elements.append(
        Paragraph("Blood Stock Details", heading)
    )

    data = [["Blood Group", "Units", "Status"]]

    stocks = BloodStock.objects.order_by("blood_group")

    for stock in stocks:

        if stock.units > 20:
            status = "Good"

        elif stock.units >= 10:
            status = "Medium"

        else:
            status = "Low"

        data.append([
            stock.blood_group,
            str(stock.units),
            status
        ])

    table = Table(data, colWidths=[150, 120, 120])

    style = [

        ("BACKGROUND", (0,0), (-1,0), colors.darkred),

        ("TEXTCOLOR", (0,0), (-1,0), colors.white),

        ("GRID", (0,0), (-1,-1), 1, colors.grey),

        ("ALIGN", (0,0), (-1,-1), "CENTER"),

        ("BOTTOMPADDING", (0,0), (-1,0), 10),

    ]

    for row in range(1, len(data)):

        if row % 2 == 0:

            style.append(

                ("BACKGROUND", (0,row), (-1,row), colors.beige)

            )

        else:

            style.append(

                ("BACKGROUND", (0,row), (-1,row), colors.whitesmoke)

            )

    table.setStyle(TableStyle(style))

    elements.append(table)

    doc.build(elements)

    return response




def export_report_excel(request):

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = (
        'attachment; filename="LifeServe_Report.xlsx"'
    )

    wb = Workbook()

    ws = wb.active
    ws.title = "Dashboard Report"

    
    # Styles
    

    title_font = Font(size=18, bold=True)

    heading_font = Font(size=12, bold=True, color="FFFFFF")

    header_fill = PatternFill(
        fill_type="solid",
        start_color="C00000",
        end_color="C00000"
    )

    center = Alignment(horizontal="center")

    
    # Title


    ws.merge_cells("A1:B1")

    ws["A1"] = "LifeServe Blood Bank Management System"

    ws["A1"].font = title_font

    ws["A1"].alignment = center

    ws["A2"] = "Generated On"

    ws["B2"] = timezone.now().strftime("%d-%m-%Y %I:%M %p")

    # ----------------------
    # Summary
    # ----------------------

    ws["A4"] = "Metric"
    ws["B4"] = "Count"

    for cell in ("A4", "B4"):
        ws[cell].font = heading_font
        ws[cell].fill = header_fill
        ws[cell].alignment = center

    summary = [

        ("Total Donors", Donor.objects.count()),

        ("Total Patients", Patient.objects.count()),

        ("Blood Requests", BloodRequest.objects.count()),

        ("Donation Requests", DonationRequest.objects.count()),

    ]

    row = 5

    for item in summary:

        ws.cell(row=row, column=1).value = item[0]

        ws.cell(row=row, column=2).value = item[1]

        row += 1

    
    # Most Requested
    

    most_requested = (
        BloodRequest.objects
        .values("blood_group")
        .annotate(total=Count("id"))
        .order_by("-total")
        .first()
    )

    row += 1

    ws.cell(row=row, column=1).value = "Most Requested Blood Group"

    if most_requested:

        ws.cell(row=row, column=2).value = (
            f"{most_requested['blood_group']} "
            f"({most_requested['total']} Requests)"
        )


    # Most Active Donor
    

    donor = (
        DonationRequest.objects
        .filter(status="Approved")
        .values(
            "donor__user__first_name",
            "donor__user__last_name"
        )
        .annotate(total=Count("id"))
        .order_by("-total")
        .first()
    )

    row += 1

    ws.cell(row=row, column=1).value = "Most Active Donor"

    if donor:

        ws.cell(row=row, column=2).value = (
            f"{donor['donor__user__first_name']} "
            f"{donor['donor__user__last_name']} "
            f"({donor['total']} Donations)"
        )

    
    # Blood Stock Table
    

    row += 3

    ws.cell(row=row, column=1).value = "Blood Group"

    ws.cell(row=row, column=2).value = "Units"

    ws.cell(row=row, column=3).value = "Status"

    for col in range(1, 4):

        cell = ws.cell(row=row, column=col)

        cell.font = heading_font

        cell.fill = header_fill

        cell.alignment = center

    row += 1

    for stock in BloodStock.objects.order_by("blood_group"):

        if stock.units > 20:
            status = "Good"

        elif stock.units >= 10:
            status = "Medium"

        else:
            status = "Low"

        ws.cell(row=row, column=1).value = stock.blood_group
        ws.cell(row=row, column=2).value = stock.units
        ws.cell(row=row, column=3).value = status

        row += 1

    
    # Column Width
    

    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 20

    wb.save(response)

    return response




def print_report(request):

    blood_stock = BloodStock.objects.all().order_by("blood_group")

    low_stock = BloodStock.objects.filter(units__lt=10)

    most_requested = (
        BloodRequest.objects
        .values("blood_group")
        .annotate(total=Count("id"))
        .order_by("-total")
        .first()
    )

    most_active_donor = (
        DonationRequest.objects
        .filter(status="Approved")
        .values(
            "donor__user__first_name",
            "donor__user__last_name"
        )
        .annotate(total=Count("id"))
        .order_by("-total")
        .first()
    )

    context = {

        "total_donors": Donor.objects.count(),

        "total_patients": Patient.objects.count(),

        "total_blood_requests": BloodRequest.objects.count(),

        "total_donation_requests": DonationRequest.objects.count(),

        "blood_stock": blood_stock,

        "low_stock": low_stock,

        "most_requested": most_requested,

        "most_active_donor": most_active_donor,

    }

    return render(
        request,
        "adminpanel/print_report.html",
        context
    )


@login_required
def admin_notifications(request):

    # print("Request User:", request.user.id, request.user.username)

    notifications = Notification.objects.filter(
        user_id=request.user.id
    ).order_by("-created_at")

    # print("Notifications:", list(notifications.values()))

    Notification.objects.filter(
        user_id=request.user.id,
        is_read=False
    ).update(is_read=True)

    return render(
        request,
        "adminpanel/admin_notifications.html",
        {
            "notifications": notifications,
        }
    )